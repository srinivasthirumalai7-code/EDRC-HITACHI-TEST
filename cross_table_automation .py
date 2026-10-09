# =========================================================
# INSTALL REQUIRED MODULES
# =========================================================
# pip install pyautogui openpyxl pywinauto keyboard pillow
# pip install opencv-python numpy
#
# =========================================================
# FINAL CROSS TABLE AUTOMATION
# =========================================================
#
# SAME SIGNAL ROUTE:
# -> COLOR CHANGE METHOD (checks if aspect reverts to RED)
#
# DIFFERENT SIGNAL ROUTE:
# -> COLOR CHANGE METHOD
#
# SHUNT SIGNALS:
# -> ORB FEATURE MATCHING (SAME LOGIC AS SHUNT CLEARANCE TOOL)
#
# =========================================================

import tkinter as tk
from tkinter import filedialog, messagebox, ttk
import pyautogui
import keyboard
import time
import threading
import os
import numpy as np
import cv2

from openpyxl import load_workbook
from pywinauto import Desktop

# =========================================================
# SETTINGS
# =========================================================

pyautogui.FAILSAFE = True

# =========================================================
# GLOBALS
# =========================================================

excel_file = ""

signal_coordinates = {}

signal_aspects = {}

shunt_signals = {}

coordinate_excel = ""
active_calling_on_track = None
running = False

paused = False

wb = None

# =========================================================
# SHUNT SNAPSHOTS FOLDER + ORB DETECTOR
# =========================================================

SHUNT_SNAPSHOTS_FOLDER = "shunt_snapshots"

if not os.path.exists(SHUNT_SNAPSHOTS_FOLDER):
    os.makedirs(SHUNT_SNAPSHOTS_FOLDER)

orb = cv2.ORB_create(nfeatures=500)

# =========================================================
# LOG FUNCTION
# =========================================================

def log(msg):

    log_box.insert(tk.END, msg + "\n")
    log_box.see(tk.END)

    root.update()

# =========================================================
# WAIT IF PAUSED
# =========================================================

def wait_if_paused():

    global paused
    global running

    while paused and running:

        time.sleep(0.2)

# =========================================================
# UPLOAD EXCEL
# =========================================================

def upload_excel():

    global excel_file

    file = filedialog.askopenfilename(
        filetypes=[("Excel Files", "*.xlsx")]
    )

    if file:

        excel_file = file

        try:
            excel_label.config(
                text=os.path.basename(file)
            )
        except:
            pass

        log(f"Excel Loaded -> {file}")
        tree.delete(*tree.get_children())

        signals = load_signals()

        for i, signal in enumerate(signals, start=1):

            total_routes = 0

            wb_temp = load_workbook(excel_file)
            ws_temp = wb_temp["CROSS TABLE"]

            for col in range(2, ws_temp.max_column + 1):

                route = ws_temp.cell(row=1, column=col).value

                if route:

                    if str(route).split("_")[0] == signal:
                        total_routes += 1

            wb_temp.close()

            tree.insert(
                "",
                "end",
                values=(
                    i,
                    signal,
                    total_routes
                )
            )
# =========================================================
# LOAD SIGNALS
# =========================================================

def load_signals():

    global excel_file

    if not excel_file:
        messagebox.showerror(
            "Error",
            "Please upload CROSS TABLE Excel first."
        )
        return []

    if not os.path.exists(excel_file):
        messagebox.showerror(
            "Error",
            f"File not found:\n{excel_file}"
        )
        return []

    if not excel_file.lower().endswith(".xlsx"):
        messagebox.showerror(
            "Error",
            "Only .xlsx files are supported."
        )
        return []

    try:
        wb = load_workbook(excel_file)
    except Exception as e:
        messagebox.showerror(
            "Excel Error",
            str(e)
        )
        return []

    ws = wb["CROSS TABLE"]

    ws = wb["CROSS TABLE"]

    signals = []

    recorded = set()

    for row in range(2, ws.max_row + 1):

        value = ws.cell(row=row, column=1).value

        if not value:
            continue

        value = str(value).strip()

        signal_name = value.split("_")[0].strip()

        if signal_name.startswith("SH"):

            if signal_name not in recorded:
                recorded.add(signal_name)
                signals.append(signal_name)

        else:

            if signal_name not in recorded:
                recorded.add(signal_name)
                signals.append(signal_name)

    wb.close()

    return signals

# =========================================================
# GET SIGNAL FROM ROUTE
# =========================================================

def get_signal_from_route(route):

    digits = ""

    for ch in route:

        if ch.isdigit():
            digits += ch
        else:
            break

    if digits == "":
        return None

    return f"S{digits}"

# =========================================================
# GET PIXEL COLOR
# =========================================================

def get_pixel_color(x, y):

    img = pyautogui.screenshot(
        region=(x, y, 1, 1)
    )

    return img.getpixel((0, 0))

# =========================================================
# COLOR CHANGE CHECK
# =========================================================

def color_changed(before, after):

    return before != after


# =========================================================
# SHUNT FEATURE COMPARISON (ORB - FROM SHUNT CLEARANCE TOOL)
# =========================================================
# Same approach as the standalone Shunt Clearance tool:
# detects key visual FEATURES and matches them, rather than
# comparing raw pixel colors. Much more robust to lighting/
# rendering variation than plain pixel comparison.
# =========================================================

def compare_snapshot_features(point, snapshot_path):

    """
    Compares the saved initial snapshot at snapshot_path
    against a freshly captured crop at `point`, using ORB
    feature matching.

    Returns (match_score_percent, details):
    - HIGH match score  -> image looks the same as before (no change)
    - LOW match score   -> a visual change was detected
    """

    try:

        if point is None or not snapshot_path:
            return 100, "Missing coordinate/snapshot"

        initial_img = cv2.imread(snapshot_path)

        if initial_img is None:
            return 100, "Cannot load initial snapshot"

        screenshot = pyautogui.screenshot()

        x, y = point

        current_pil = screenshot.crop(
            (x - 50, y - 50, x + 50, y + 50)
        )

        current_img = cv2.cvtColor(
            np.array(current_pil),
            cv2.COLOR_RGB2BGR
        )

        initial_gray = cv2.cvtColor(initial_img, cv2.COLOR_BGR2GRAY)
        current_gray = cv2.cvtColor(current_img, cv2.COLOR_BGR2GRAY)

        kp1, des1 = orb.detectAndCompute(initial_gray, None)
        kp2, des2 = orb.detectAndCompute(current_gray, None)

        if des1 is None or des2 is None:
            return 0, "Cannot detect features"

        bf = cv2.BFMatcher(cv2.NORM_HAMMING, crossCheck=True)

        matches = bf.match(des1, des2)

        if len(matches) > 0:

            good_matches = [m for m in matches if m.distance < 50]

            match_score = (
                (len(good_matches) / max(len(kp1), len(kp2))) * 100
                if max(len(kp1), len(kp2)) > 0 else 0
            )

        else:

            match_score = 0

        return match_score, f"Features matched: {len(matches)}"

    except Exception as e:

        log(f"FEATURE COMPARISON ERROR -> {e}")

        return 0, str(e)


def shunt_signal_changed(shunt_name):

    """
    Uses ORB feature matching (Shunt Clearance tool logic)
    to decide if this shunt signal actually changed.

    Checks BOTH:
    - ASPECT INDICATOR: the real visual check - did the
      shunt route actually set/clear.
    - INITIATION INDICATION: a separate check that catches
      the case where the route gets initiated/accepted by
      the interlocking but the aspect itself never visibly
      changes - same purpose as INITIATION_INDICATION for
      MAIN/CALLING_ON signals.

    If EITHER shows a change, the result is NOT LOCKED.
    Only if BOTH stay unchanged is the result LOCKED.

    Returns True  -> a change was detected  (NOT LOCKED)
    Returns False -> nothing changed        (LOCKED)
    """

    if shunt_name not in shunt_signals:
        return False

    info = shunt_signals[shunt_name]

    # Give the interlocking time to settle before comparing,
    # same as the Shunt Clearance tool's proven 7 second wait.
    smart_sleep(7)

    changed = False

    # =====================================
    # ASPECT INDICATOR
    # =====================================

    aspect_point = info.get("aspect_indicator")
    aspect_snapshot = info.get("aspect_snapshot")

    if aspect_point and aspect_snapshot:

        match_score, details = compare_snapshot_features(
            aspect_point,
            aspect_snapshot
        )

        log(
            f"{shunt_name} ASPECT INDICATOR MATCH -> "
            f"{match_score:.2f}% ({details})"
        )

        if match_score < 70:

            changed = True

        elif match_score > 85:

            pass

        else:

            # Ambiguous - benefit of the doubt goes to
            # flagging a possible failure, not missing one
            changed = True

    # =====================================
    # INITIATION INDICATION
    # =====================================

    init_point = info.get("initiation_indicator")
    initiation_snapshot = info.get("initiation_snapshot")

    if init_point and initiation_snapshot:

        match_score, details = compare_snapshot_features(
            init_point,
            initiation_snapshot
        )

        log(
            f"{shunt_name} INITIATION INDICATION MATCH -> "
            f"{match_score:.2f}% ({details})"
        )

        if match_score < 70:

            changed = True

        elif match_score > 85:

            pass

        else:

            changed = True

    elif init_point:

        # No dedicated snapshot for this point (e.g. loaded from the
        # universal coordinate file, which only stores one snapshot
        # per shunt signal, tied to the aspect indicator). Route
        # Initiation Indicator doesn't need ORB matching anyway - it's
        # a simple lamp: lit yellow means the route was accepted by
        # the interlocking, same convention used everywhere else in
        # the suite.

        r, g, b = get_pixel_color(init_point[0], init_point[1])

        initiated = (r > 150 and g > 150)

        log(
            f"{shunt_name} INITIATION INDICATION COLOR -> "
            f"RGB({r},{g},{b}) "
            f"{'YELLOW (route accepted)' if initiated else 'not yellow'}"
        )

        if initiated:

            changed = True

    return changed


# =========================================================
# CALLING ON CHECK
# =========================================================

def is_calling_on_signal(signal_name):

    signal_name = signal_name.upper()

    first_part = signal_name.split("_")[0]

    if first_part.endswith("C"):
        return True

    return False


# =========================================================
# GET TRACK FOR CALLING ON SIGNAL
# =========================================================

def get_calling_on_track(signal_name, tracks_ws):

    signal_name = str(signal_name).strip().upper()

    # 1C_A -> 1C
    # 1C_B -> 1C
    # 32C_A -> 32C
    # 32C_B -> 32C

    if "_" in signal_name:

        signal_name = signal_name.split("_")[0]
    log(f"SEARCHING TRACK FOR -> {signal_name}")

    for row in range(2, tracks_ws.max_row + 1):

        sig = tracks_ws.cell(row=row, column=1).value
        trk = tracks_ws.cell(row=row, column=2).value

        log(f"EXCEL SIGNAL = [{sig}]")
        log(f"SEARCH SIGNAL = [{signal_name}]")

        if not sig or not trk:
            continue

        sig = str(sig).strip().upper()

        log(f"AFTER STRIP = [{sig}]")

        if sig == signal_name:
            log(f"MATCH FOUND -> {sig}")
            log(f"TRACK FOUND -> {trk}")

            return str(trk).strip()

    log(f"NO TRACK FOUND FOR -> {signal_name}")

    return None
# =========================================================
# TRACK DOWN
# =========================================================
def track_down(track_name):

    log("TRACK_DOWN FUNCTION CALLED")
    log(f"TRACK NAME -> {track_name}")

    log(f"TRACK DOWN -> {track_name}")


    keyboard.press_and_release("ctrl+b")
    smart_sleep(2)

    if not click_control("50051"):
        log("50051 NOT FOUND")
        return False

    smart_sleep(1)

    if not click_control("OK"):
        log("OK NOT FOUND")
        return False

    smart_sleep(1)

    if not click_control(track_name):
        log(f"TRACK NOT FOUND -> {track_name}")
        return False

    smart_sleep(1)

    if not click_control("Transmit"):
        log("Transmit NOT FOUND")
        return False

    smart_sleep(1)

    click_control("OK")

    smart_sleep(1)

    click_control("Cancel")

    smart_sleep(1)

    log("TRACK DOWN COMPLETE")
    return True
def track_up(track_name):

    log(f"TRACK UP -> {track_name}")

    keyboard.press_and_release("ctrl+b")
    smart_sleep(2)

    if not click_control("50051"):
        log("50051 NOT FOUND")
        return False

    smart_sleep(1)

    if not click_control("OK"):
        log("OK NOT FOUND")
        return False

    smart_sleep(1)

    if not click_control(track_name):
        log(f"TRACK NOT FOUND -> {track_name}")
        return False

    smart_sleep(1)

    if not click_control("Transmit"):
        log("Transmit NOT FOUND")
        return False

    smart_sleep(1)

    click_control("OK")

    smart_sleep(1)

    click_control("Cancel")

    smart_sleep(1)

    log("TRACK UP COMPLETE")
    return True

# =========================================================
# GET DROPDOWN ITEMS
# =========================================================

def get_dropdown_items():

    items = []

    try:

        menu = Desktop(
            backend="uia"
        ).window(control_type="Menu")

        for item in menu.descendants(
            control_type="MenuItem"
        ):

            text = item.window_text().strip()

            if text:
                items.append((text, item))

    except Exception as e:

        log(str(e))

    return items

def try_select_route(route):

    try:

        menu = Desktop(
            backend="uia"
        ).window(control_type="Menu")

        item = menu.child_window(
            title=route,
            control_type="MenuItem"
        )

        if item.exists(timeout=0.5):

            item.click_input()

            return True

        log(f"Menu Item Not Found -> {route}")

        return False

    except Exception as e:

        log(str(e))

        return False
# =========================================================
# CONVERT SHUNT ROUTE NAME FOR MENU
# =========================================================

def get_actual_menu_route(route):

    route = route.strip()

    # SH9_A -> 9_A
    # SH9_B -> 9_B
    # SH21_A -> 21_A
    # SH21_B -> 21_B

    if route.upper().startswith("SH"):

        return route[2:]

    return route
# =========================================================
# CLICK SPECIAL MENU
# =========================================================


def click_special_menu(menu_name):

    try:

        menu = Desktop(
            backend="uia"
        ).window(control_type="Menu")

        for item in menu.descendants():

            try:

                if item.window_text().strip() == menu_name:

                    item.click_input()

                    log(f"Clicked -> {menu_name}")

                    smart_sleep(0.3)

                    return True

            except:
                pass

        log(f"{menu_name} Not Found")

        return False

    except Exception as e:

        log(str(e))

        return False
# =========================================================
# CTRL+B WINDOW DUMP
# =========================================================

def dump_ctrl_b_window():

    try:

        windows = Desktop(backend="uia").windows()

        for dlg in windows:

            log("================================")
            log(f"WINDOW : {dlg.window_text()}")

            for item in dlg.descendants():

                try:

                    log(
                        f"{item.window_text()} -> "
                        f"{item.element_info.control_type}"
                    )

                except:
                    pass

    except Exception as e:

        log(str(e))

# =========================================================
# CLICK CONTROL (CTRL+B WINDOW)
# =========================================================
def click_control(control_name):

    try:

        # =====================================
        # FAST PATH - CHECK TOPMOST WINDOW FIRST
        # -------------------------------------
        # The CTRL+B dialog is almost always the
        # active/topmost window, so checking it
        # first avoids scanning every other open
        # window and its full control tree.
        # =====================================

        try:

            top = Desktop(backend="uia").top_window()

            for item in top.descendants():

                try:

                    txt = item.window_text().strip()

                    if txt == control_name:

                        log(
                            f"FOUND (FAST) -> {control_name} "
                            f"({item.element_info.control_type})"
                        )

                        item.click_input()

                        return True

                except:
                    pass

        except Exception as e:

            log(f"TOP WINDOW CHECK FAILED -> {e}")

        # =====================================
        # FALLBACK - SCAN ALL WINDOWS
        # =====================================

        windows = Desktop(backend="uia").windows()

        for dlg in windows:

            try:

                for item in dlg.descendants():

                    try:

                        txt = item.window_text().strip()

                        if txt == control_name:

                            log(
                                f"FOUND -> {control_name} "
                                f"({item.element_info.control_type})"
                            )

                            item.click_input()

                            return True

                    except:
                        pass

            except:
                pass

        log(f"NOT FOUND -> {control_name}")

        return False

    except Exception as e:

        log(str(e))

        return False


# =========================================================
# WAIT IF PAUSED
# =========================================================

def check_pause():

    global paused
    global running

    while paused and running:

        time.sleep(0.2)
def release_signal(signal_name):

    if signal_name not in signal_coordinates:
        return

    x, y = signal_coordinates[signal_name]

    log(f"{signal_name} SIGNAL CANCEL")

    pyautogui.click(x, y)

    smart_sleep(1)

    click_special_menu("Signal Cancel")

    smart_sleep(1)

    log(f"{signal_name} ROUTE RELEASE")

    pyautogui.click(x, y)

    smart_sleep(1)

    click_special_menu("Route Release")

    smart_sleep(8)
# =========================================================
# START AUTOMATION
# =========================================================

def start_automation():

    status_label.config(
        text="RUNNING",
        fg="#16a34a"
    )

    root.iconify()

    # Wait 2 seconds after minimizing
    root.after(
        2000,
        lambda: threading.Thread(
            target=automation,
            daemon=True
        ).start()
    )
# =========================================================
# MAIN AUTOMATION
# =========================================================

def automation():

    global running
    global paused
    global wb

    if excel_file == "":

        messagebox.showerror(
            "Error",
            "Upload Excel First"
        )

        return

    running = True
    global active_calling_on_track

    active_calling_on_track = None
    wb = load_workbook(excel_file)

    ws = wb["CROSS TABLE"]

    tracks_ws = wb["TRACKS"]
    if not signal_coordinates:
        messagebox.showerror(
            "Missing Data",
            "Signal coordinates not loaded.\nPlease record or upload coordinate Excel."
        )
        return

    # =====================================================
    # MAP EACH ROUTE NAME -> ITS OWN ROW IN COLUMN A
    # =====================================================
    # This is required because multiple routes can belong
    # to the same signal (e.g. 1_A, 1_A1, 1_B all belong to
    # signal "1"), and each one has its OWN row in the
    # CROSS TABLE that results must be written into.
    # =====================================================

    route_row_map = {}

    for r in range(2, ws.max_row + 1):

        route_label = ws.cell(row=r, column=1).value

        if route_label:

            route_row_map[str(route_label).strip()] = r
    # =====================================================
    # GET ALL ROUTES
    # =====================================================

    routes = []

    log("LOADING ROUTES FROM EXCEL...")

    for col in range(2, ws.max_column + 1):

        value = ws.cell(
            row=1,
            column=col
        ).value

        if value:
            value = str(value).strip()

            routes.append(value)

            log(f"ROUTE FOUND -> {value}")

    if not routes:
        messagebox.showerror(
            "Error",
            "No routes found in CROSS TABLE sheet"
        )

        return
    # =====================================================
    # PROCESS SIGNALS
    # =====================================================
    processed_signals = set()
    for row in range(2, ws.max_row + 1):

        check_pause()

        if not running:
            break

        base_signal = ws.cell(
            row=row,
            column=1
        ).value

        if not base_signal:
            continue

        base_signal = str(base_signal).strip()

        base_signal_no = base_signal.split("_")[0]
        if base_signal_no in processed_signals:
            continue

        processed_signals.add(base_signal_no)
        log(f"Base Signal : {base_signal}")
        log(f"Coordinate Key : {base_signal_no}")
        if base_signal_no not in signal_coordinates:
            continue

        log("")
        log("================================")
        log(f"PROCESSING SIGNAL -> {base_signal}")
        log("================================")

        # =================================================
        # GET ALL ROUTES OF THIS SIGNAL
        # =================================================

        signal_routes = []

        for route in routes:

            if route.split("_")[0] == base_signal_no:
                signal_routes.append(route)

        if not signal_routes:
            continue

        # =================================================
        # PROCESS EACH ROUTE
        # =================================================

        for base_route in signal_routes:

            check_pause()

            if not running:
                break

            # =========================================
            # FIND THIS ROUTE'S OWN ROW TO WRITE RESULTS
            # =========================================

            target_row = route_row_map.get(base_route)

            if target_row is None:

                log(f"ROW NOT FOUND FOR ROUTE -> {base_route}")

                continue

            log("")
            log(f"SETTING ROUTE -> {base_route}")
            calling_on_route = is_calling_on_signal(base_route)

            log(f"ROUTE = {base_route}")
            log(f"CALLING ON = {calling_on_route}")
            log(f"CHECKING CALLING ON -> {base_route}")
            log(f"RESULT -> {is_calling_on_signal(base_route)}")
            if calling_on_route:
                log("ENTERED CALLING ON BLOCK")
                track_name = get_calling_on_track(
                    base_route,
                    tracks_ws
                )

                if track_name:

                    if active_calling_on_track is None:

                        track_down(track_name)
                        active_calling_on_track = track_name

                    elif active_calling_on_track != track_name:

                        log(f"TRACK UP -> {active_calling_on_track}")

                        track_up(active_calling_on_track)

                        smart_sleep(1)

                        track_down(track_name)

                        active_calling_on_track = track_name

                    else:

                        log(f"SAME TRACK ACTIVE -> {track_name}")
                if track_name:

                    # Same track already down
                    if active_calling_on_track == track_name:

                        log(f"TRACK ALREADY DOWN -> {track_name}")

                    else:

                        # Previous track restore
                        if active_calling_on_track:
                            log(f"TRACK UP -> {active_calling_on_track}")

                            track_up(active_calling_on_track)

                            smart_sleep(1)

                        log("STARTING TRACK DOWN")

                        track_down(track_name)

                        active_calling_on_track = track_name

                        log("TRACK DOWN COMPLETED")

                else:

                    log("TRACK NOT FOUND")
            if is_calling_on_signal(base_route):

                base_coordinate_name = base_signal_no

            else:

                base_coordinate_name = base_signal_no

            bx, by = signal_coordinates[base_coordinate_name]

            # =============================================
            # CAPTURE BASE SIGNAL RED COLOR (BEFORE SETTING)
            # =============================================

            base_is_shunt = base_signal_no.startswith("SH")

            base_red_before = None

            if (
                not calling_on_route
                and not base_is_shunt
                and base_coordinate_name in signal_aspects
                and "RED" in signal_aspects[base_coordinate_name]
            ):

                brx, bry = signal_aspects[base_coordinate_name]["RED"]

                base_red_before = get_pixel_color(brx, bry)

            # =============================================
            # CAPTURE CALLING ON LAMP COLOR (BEFORE SETTING)
            # =============================================

            base_callingon_before = None

            if (
                calling_on_route
                and base_coordinate_name in signal_aspects
                and "CALLING_ON" in signal_aspects[base_coordinate_name]
            ):

                cox, coy = signal_aspects[base_coordinate_name]["CALLING_ON"]

                base_callingon_before = get_pixel_color(cox, coy)

            # =============================================
            # OPEN MENU
            # =============================================

            pyautogui.click(bx, by)
            check_pause()

            smart_sleep(1)
            check_pause()

            # =============================================
            # SET ROUTE
            # =============================================

            menu_route = get_actual_menu_route(base_route)

            selected = try_select_route(menu_route)
            check_pause()

            if not selected:

                log(f"FAILED -> {base_route}")

                pyautogui.press("esc")

                continue

            log(f"SET -> {base_route}")

            smart_sleep(4)
            check_pause()

            # =============================================
            # VERIFY BASE SIGNAL CLEARED (RED -> YELLOW/GREEN)
            # =============================================

            if base_red_before is not None:

                brx, bry = signal_aspects[base_coordinate_name]["RED"]

                base_red_after = get_pixel_color(brx, bry)

                if not color_changed(base_red_before, base_red_after):

                    log(f"BASE SIGNAL DID NOT CLEAR (STILL RED) -> {base_route}")

                    log("MARKING ALL ROUTES AS FAILED FOR THIS ROW")

                    # =====================================
                    # MARK EVERY COLUMN IN THIS ROW FAILED
                    # =====================================

                    for fcol in range(2, ws.max_column + 1):

                        fail_test_route = ws.cell(
                            row=1,
                            column=fcol
                        ).value

                        if not fail_test_route:
                            continue

                        fail_test_route = str(fail_test_route).strip()

                        if fail_test_route == base_route:

                            ws.cell(
                                row=target_row,
                                column=fcol
                            ).value = "-"

                        else:

                            ws.cell(
                                row=target_row,
                                column=fcol
                            ).value = "FAILED"

                    pyautogui.press("esc")

                    smart_sleep(0.5)
                    check_pause()

                    # =====================================
                    # SIGNAL CANCEL (BASE FAILED TO CLEAR)
                    # =====================================

                    log("SIGNAL CANCEL (BASE FAILED)")

                    pyautogui.click(bx, by)
                    check_pause()

                    smart_sleep(0.5)
                    check_pause()

                    try_select_route("Signal Cancel")
                    check_pause()

                    smart_sleep(1)
                    check_pause()

                    # =====================================
                    # ROUTE RELEASE (BASE FAILED TO CLEAR)
                    # =====================================

                    log("ROUTE RELEASE (BASE FAILED)")

                    pyautogui.click(bx, by)
                    check_pause()

                    smart_sleep(0.5)
                    check_pause()

                    try_select_route("Route Release")
                    check_pause()

                    smart_sleep(8)
                    check_pause()

                    log(f"{base_route} FAILED -> MOVING TO NEXT ROUTE")

                    continue

                log(f"BASE SIGNAL CLEARED (RED -> YELLOW/GREEN) -> {base_route}")

            # =============================================
            # VERIFY CALLING ON LAMP LIT (BEFORE TESTING)
            # =============================================

            if base_callingon_before is not None:

                cox, coy = signal_aspects[base_coordinate_name]["CALLING_ON"]

                base_callingon_after = get_pixel_color(cox, coy)

                if not color_changed(base_callingon_before, base_callingon_after):

                    log(f"CALLING ON LAMP DID NOT LIGHT -> {base_route}")

                    log("MARKING ALL ROUTES AS FAILED FOR THIS ROW")

                    # =====================================
                    # MARK EVERY COLUMN IN THIS ROW FAILED
                    # =====================================

                    for fcol in range(2, ws.max_column + 1):

                        fail_test_route = ws.cell(
                            row=1,
                            column=fcol
                        ).value

                        if not fail_test_route:
                            continue

                        fail_test_route = str(fail_test_route).strip()

                        if fail_test_route == base_route:

                            ws.cell(
                                row=target_row,
                                column=fcol
                            ).value = "-"

                        else:

                            ws.cell(
                                row=target_row,
                                column=fcol
                            ).value = "FAILED"

                    pyautogui.press("esc")

                    smart_sleep(0.5)
                    check_pause()

                    # =====================================
                    # SIGNAL CANCEL (CALLING ON FAILED)
                    # =====================================

                    log("SIGNAL CANCEL (CALLING ON FAILED)")

                    pyautogui.click(bx, by)
                    check_pause()

                    smart_sleep(0.5)
                    check_pause()

                    try_select_route("Signal Cancel")
                    check_pause()

                    smart_sleep(1)
                    check_pause()

                    # =====================================
                    # ROUTE RELEASE (CALLING ON FAILED)
                    # =====================================

                    log("ROUTE RELEASE (CALLING ON FAILED)")

                    pyautogui.click(bx, by)
                    check_pause()

                    smart_sleep(0.5)
                    check_pause()

                    try_select_route("Route Release")
                    check_pause()

                    smart_sleep(8)
                    check_pause()

                    # =====================================
                    # TRACK UP (RESTORE TRACK DROPPED
                    # FOR THIS FAILED CALLING ON ROUTE)
                    # =====================================

                    if active_calling_on_track:

                        log(f"TRACK UP -> {active_calling_on_track}")

                        track_up(active_calling_on_track)

                        smart_sleep(1)
                        check_pause()

                        active_calling_on_track = None

                    log(f"{base_route} FAILED -> MOVING TO NEXT ROUTE")

                    continue

                log(f"CALLING ON LAMP LIT -> {base_route}")

            # =============================================
            # TEST ALL ROUTES
            # =============================================

            current_test_track = None
            for col in range(2, ws.max_column + 1):

                check_pause()

                if not running:
                    break

                test_route = ws.cell(
                    row=1,
                    column=col
                ).value

                if not test_route:
                    continue

                test_route = str(test_route).strip()

                if test_route == base_route:
                    ws.cell(
                        row=target_row,
                        column=col
                    ).value = "-"

                    continue

                target_signal = test_route.split("_")[0]

                if target_signal not in signal_coordinates:
                    continue

                tx, ty = signal_coordinates[target_signal]

                log("")
                log(f"Testing -> {test_route}")
                # Leaving a calling-on group

                if (
                        current_test_track
                        and
                        not is_calling_on_signal(test_route)
                ):
                    log(f"TRACK UP -> {current_test_track}")

                    track_up(current_test_track)

                    smart_sleep(1)

                    current_test_track = None
                # =========================================
                # CALLING ON TEST ROUTE TRACK DOWN
                # =========================================

                if is_calling_on_signal(test_route):

                    log(f"CALLING ON TEST ROUTE -> {test_route}")

                    track_name = get_calling_on_track(
                        test_route,
                        tracks_ws
                    )

                    if track_name:

                        if current_test_track != track_name:

                            if current_test_track:
                                log(f"TRACK UP -> {current_test_track}")
                                track_up(current_test_track)
                                smart_sleep(1)

                            log(f"TRACK DOWN FOR TEST ROUTE -> {track_name}")

                            track_down(track_name)

                            current_test_track = track_name

                            smart_sleep(1)

                    else:

                        log(f"TRACK NOT FOUND FOR -> {test_route}")
                # =========================================
                # SAME SIGNAL
                # =========================================

                if target_signal == base_signal_no:

                    # =====================================
                    # SAME SIGNAL - ALTERNATE ROUTE TEST
                    # -----------------------------------
                    # Base route is already SET. We now try
                    # to select a DIFFERENT route of the
                    # SAME signal.
                    # =====================================

                    if target_signal.startswith("SH"):

                        # =================================
                        # SHUNT - SAME SIGNAL, ALT ROUTE
                        # ---------------------------------
                        # Uses ORB feature matching (Shunt
                        # Clearance tool logic) instead of
                        # the RED/aspect check, since shunt
                        # signals don't have RED/YELLOW/
                        # GREEN aspects.
                        # =================================

                        pyautogui.click(tx, ty)

                        check_pause()

                        smart_sleep(1)

                        menu_route = get_actual_menu_route(test_route)

                        selected = try_select_route(menu_route)

                        check_pause()

                        if shunt_signal_changed(target_signal):

                            log(f"{test_route} SHUNT CHANGED -> FAIL")

                            result = "NOT LOCKED"

                            release_signal(target_signal)

                        else:

                            log(f"{test_route} SHUNT UNCHANGED -> PASS")

                            result = "LOCKED"

                    else:

                        # =================================
                        # MAIN / CALLING ON - SAME SIGNAL
                        # ---------------------------------
                        # This signal is currently YELLOW/
                        # GREEN (RED lamp OFF). If RED turns
                        # back ON, that's NOT LOCKED. If only
                        # INITIATION INDICATION lights without
                        # RED reappearing, that's FAILED.
                        # =================================

                        before_colors = {}

                        if target_signal in signal_aspects:

                            for aspect_name in ("RED", "INITIATION_INDICATION"):

                                if aspect_name in signal_aspects[target_signal]:

                                    ax, ay = signal_aspects[target_signal][aspect_name]

                                    before_colors[aspect_name] = get_pixel_color(ax, ay)

                        pyautogui.click(tx, ty)

                        check_pause()

                        smart_sleep(1)

                        menu_route = get_actual_menu_route(test_route)

                        selected = try_select_route(menu_route)

                        check_pause()

                        # Close any stuck-open menu so it can't
                        # cover the aspect lamps we're about to read.
                        pyautogui.press("esc")

                        check_pause()

                        # =============================
                        # WAIT 6 SECONDS, THEN CHECK ONCE
                        # =============================

                        smart_sleep(6)
                        check_pause()

                        aspect_changed = False
                        initiation_changed = False

                        if target_signal in signal_aspects:

                            if "RED" in signal_aspects[target_signal]:

                                rx, ry = signal_aspects[target_signal]["RED"]

                                before_red = before_colors["RED"]

                                after_red = get_pixel_color(rx, ry)

                                if color_changed(before_red, after_red):
                                    aspect_changed = True

                            if "INITIATION_INDICATION" in signal_aspects[target_signal]:

                                iix, iiy = signal_aspects[target_signal]["INITIATION_INDICATION"]

                                before_ii = before_colors["INITIATION_INDICATION"]

                                after_ii = get_pixel_color(iix, iiy)

                                if color_changed(before_ii, after_ii):
                                    initiation_changed = True

                        if aspect_changed:

                            log(f"{test_route} CHANGED TO RED -> NOT LOCKED")

                            result = "NOT LOCKED"

                        elif initiation_changed:

                            log(f"{test_route} INITIATION SEEN, RED DID NOT RETURN -> FAILED")

                            result = "FAILED"

                        else:

                            log(f"{test_route} STAYED YELLOW/GREEN, NO INITIATION -> LOCKED")

                            result = "LOCKED"

                # =========================================
                # DIFFERENT SIGNAL
                # =========================================
                else:

                    calling_on_target = is_calling_on_signal(
                        target_signal
                    )
                    is_shunt = target_signal.startswith("SH")
                    before_colors = {}

                    if target_signal in signal_aspects:

                        for aspect_name, coord in signal_aspects[target_signal].items():

                            ax, ay = coord

                            before_colors[aspect_name] = get_pixel_color(
                                ax,
                                ay
                            )

                    pyautogui.click(tx, ty)
                    check_pause()

                    smart_sleep(1)
                    check_pause()

                    menu_route = get_actual_menu_route(test_route)

                    try_select_route(menu_route)
                    check_pause()

                    # Close any stuck-open menu so it can't
                    # cover the aspect lamps we're about to read.
                    pyautogui.press("esc")

                    check_pause()

                    # =====================================
                    # WAIT 6 SECONDS, THEN CHECK ONCE
                    # =====================================

                    smart_sleep(6)
                    check_pause()

                    aspect_changed = False
                    initiation_changed = False
                    changed_aspects = {}

                    if target_signal in signal_aspects:

                        for aspect_name, coord in signal_aspects[target_signal].items():

                            ax, ay = coord

                            after = get_pixel_color(ax, ay)

                            before = before_colors[aspect_name]

                            if color_changed(before, after):

                                changed_aspects[aspect_name] = after

                                if aspect_name == "INITIATION_INDICATION":

                                    initiation_changed = True

                                else:

                                    aspect_changed = True

                    check_pause()
                    # =========================================
                    # SHUNT SIGNAL CHECK
                    # =========================================

                    if target_signal.startswith("SH"):

                        if shunt_signal_changed(target_signal):

                            log(f"{test_route} SHUNT CHANGED -> FAIL")

                            result = "NOT LOCKED"

                            release_signal(target_signal)

                        else:

                            log(f"{test_route} SHUNT UNCHANGED -> PASS")

                            result = "LOCKED"

                        ws.cell(
                            row=target_row,
                            column=col
                        ).value = result

                        pyautogui.press("esc")

                        smart_sleep(0.5)

                        continue
                    if calling_on_target:

                        if aspect_changed:

                            result = "CONDITIONALLY LOCKED"

                            log(f"{target_signal} CALLING ON LAMP CHANGED -> {result}")

                            log(f"{target_signal} SIGNAL CANCEL")

                            pyautogui.click(tx, ty)

                            smart_sleep(1)

                            click_special_menu("Signal Cancel")

                            smart_sleep(1)

                            log(f"{target_signal} ROUTE RELEASE")

                            pyautogui.click(tx, ty)

                            smart_sleep(0.5)

                            click_special_menu("Route Release")

                            smart_sleep(8)

                        elif initiation_changed:

                            result = "FAILED"

                            log(f"{target_signal} INITIATION SEEN, LAMP DID NOT SET -> {result}")

                            log(f"{target_signal} SIGNAL CANCEL")

                            pyautogui.click(tx, ty)

                            smart_sleep(1)

                            click_special_menu("Signal Cancel")

                            smart_sleep(1)

                            log(f"{target_signal} ROUTE RELEASE")

                            pyautogui.click(tx, ty)

                            smart_sleep(0.5)

                            click_special_menu("Route Release")

                            smart_sleep(8)

                        else:

                            result = "LOCKED"

                        ws.cell(
                            row=target_row,
                            column=col
                        ).value = result

                        pyautogui.press("esc")

                        smart_sleep(0.5)

                        continue
                    if aspect_changed:

                        log(f"{test_route} ASPECT CHANGED -> NOT LOCKED")

                        result = "NOT LOCKED"

                        release_signal(target_signal)

                    elif initiation_changed:

                        log(f"{test_route} INITIATION SEEN, ASPECT DID NOT CHANGE -> FAILED")

                        result = "FAILED"

                        release_signal(target_signal)

                    else:

                        result = "LOCKED"


                # =========================================
                # WRITE RESULT
                # =========================================

                ws.cell(
                    row=target_row,
                    column=col
                ).value = result
                pyautogui.press("esc")

                smart_sleep(0.5)
                check_pause()

            # =============================================
            # SIGNAL CANCEL
            # =============================================

            log("SIGNAL CANCEL")
            if current_test_track:
                log(f"FINAL TRACK UP -> {current_test_track}")

                track_up(current_test_track)

                smart_sleep(1)

                current_test_track = None
            pyautogui.click(bx, by)
            check_pause()

            smart_sleep(0.5)
            check_pause()

            try_select_route("Signal Cancel")
            check_pause()

            smart_sleep(1)
            check_pause()

            # =============================================
            # ROUTE RELEASE
            # =============================================

            log("ROUTE RELEASE")

            pyautogui.click(bx, by)
            check_pause()

            smart_sleep(0.5)
            check_pause()

            try_select_route("Route Release")
            check_pause()

            smart_sleep(8)
            check_pause()

            log(f"ACTIVE TRACK = {active_calling_on_track}")
            log(f"{base_route} COMPLETED")
    # 👇 ADD IT HERE (IMPORTANT POSITION)

    # =====================================================
    # SAVE FILE
    # =====================================================
    output_file = filedialog.asksaveasfilename(
        defaultextension=".xlsx",
        filetypes=[("Excel Files", "*.xlsx")],
        title="Save Result Excel"
    )

    if not output_file:
        return
    wb.save(output_file)

    wb.close()

    log("")
    log(f"FILE SAVED -> {output_file}")

    messagebox.showinfo(
        "Completed",
        f"Updated File Saved:\n{output_file}"
    )
    running = False

    status_label.config(
        text="COMPLETED",
        fg="#16a34a"
    )

    root.deiconify()
# =========================================================
# STOP
# =========================================================

def stop_automation():

    global running
    global paused
    global wb

    running = False

    paused = False
    status_label.config(
        text="STOPPED",
        fg="#dc2626"
    )

    root.deiconify()
    log("")
    log("================================")
    log("STOP BUTTON PRESSED")
    log("================================")

    try:

        output_file = filedialog.asksaveasfilename(
            defaultextension=".xlsx",
            filetypes=[("Excel Files", "*.xlsx")],
            title="Save Result Excel"
        )

        if not output_file:
            log("Save Cancelled")

            return
        if wb is not None:

            wb.save(output_file)

            log(f"PARTIAL FILE SAVED -> {output_file}")

    except Exception as e:

        log(f"SAVE ERROR -> {e}")

    wb = None

    log("Automation Stopped")

def smart_sleep(seconds):

    start = time.time()

    while (time.time() - start) < seconds:

        check_pause()

        if not running:
            return

        time.sleep(0.1)
# =========================================================
# PAUSE / RESUME
# =========================================================
def pause_automation():

    keyboard_pause_toggle()



# =========================================================
# LOAD UNIVERSAL YARD COORDINATES
# =========================================================
# Reads the multi-sheet Universal Yard Coordinate file (the
# same file every other testing program in the suite loads
# from) instead of this program's own bespoke coordinate
# Excel format, and populates signal_coordinates /
# signal_aspects / shunt_signals to match what the rest of
# this program expects.
#
# NOTE 1: this program names shunt signals with a PREFIX
# ("SH9"), while the universal file (and every other program)
# uses a SUFFIX ("9SH") - names are converted on load.
#
# NOTE 2: the universal file stores only ONE snapshot per
# shunt signal, tied to its Indicator point - there's no
# second snapshot for the Initiation Indication point. That's
# fine: shunt_signal_changed() checks Initiation Indication by
# color (lit yellow = route accepted) instead of ORB matching
# when loaded this way, same convention used everywhere else
# in the suite. Only the Aspect Indicator uses ORB matching.
# =========================================================

def load_universal_coordinates():

    global signal_coordinates
    global signal_aspects
    global shunt_signals

    file = filedialog.askopenfilename(
        title="Select Universal Yard Coordinates",
        filetypes=[("Excel Files", "*.xlsx")]
    )

    if not file:
        return

    wb_u = load_workbook(file)

    signal_coordinates.clear()
    signal_aspects.clear()
    shunt_signals.clear()

    def parse_int(value):
        if value is None:
            return None
        if isinstance(value, str) and value.strip() == "":
            return None
        try:
            return int(value)
        except (TypeError, ValueError):
            return None

    def parse_point(x_val, y_val):
        x = parse_int(x_val)
        y = parse_int(y_val)
        if x is None or y is None:
            return None
        return (x, y)

    def header_index_of(ws_u):
        header_row = next(ws_u.iter_rows(min_row=1, max_row=1, values_only=True), None)
        if not header_row:
            return {}
        idx = {}
        for i, val in enumerate(header_row):
            if val is not None:
                idx[str(val).strip().upper()] = i
        return idx

    def cell(row, idx):
        return row[idx] if idx is not None and len(row) > idx else None

    loaded_main = loaded_shunt = loaded_cal = 0

    # ---- MAIN ----
    if "MAIN" in wb_u.sheetnames:

        ws_u = wb_u["MAIN"]
        hi = header_index_of(ws_u)
        sig_col = hi.get("SIGNAL")

        if sig_col is not None:

            for row in ws_u.iter_rows(min_row=2, values_only=True):

                if not row or cell(row, sig_col) is None:
                    continue

                sig = str(cell(row, sig_col)).strip().upper()

                menu = parse_point(cell(row, hi.get("MENU_X")), cell(row, hi.get("MENU_Y")))

                if menu:
                    signal_coordinates[sig] = menu

                aspects = {}

                red = parse_point(cell(row, hi.get("RED_X")), cell(row, hi.get("RED_Y")))
                yellow = parse_point(cell(row, hi.get("YELLOW_X")), cell(row, hi.get("YELLOW_Y")))
                dy = parse_point(cell(row, hi.get("DOUBLE_YELLOW_X")), cell(row, hi.get("DOUBLE_YELLOW_Y")))
                green = parse_point(cell(row, hi.get("GREEN_X")), cell(row, hi.get("GREEN_Y")))
                route_init = parse_point(cell(row, hi.get("ROUTEINIT_X")), cell(row, hi.get("ROUTEINIT_Y")))

                if red:
                    aspects["RED"] = red
                if yellow:
                    aspects["YELLOW"] = yellow
                if dy:
                    aspects["DOUBLE_YELLOW"] = dy
                if green:
                    aspects["GREEN"] = green
                if route_init:
                    aspects["INITIATION_INDICATION"] = route_init

                if aspects:
                    signal_aspects[sig] = aspects

                loaded_main += 1

    # ---- CAL ----
    if "CAL" in wb_u.sheetnames:

        ws_u = wb_u["CAL"]
        hi = header_index_of(ws_u)
        sig_col = hi.get("SIGNAL")

        if sig_col is not None:

            for row in ws_u.iter_rows(min_row=2, values_only=True):

                if not row or cell(row, sig_col) is None:
                    continue

                sig = str(cell(row, sig_col)).strip().upper()

                menu = parse_point(cell(row, hi.get("MENU_X")), cell(row, hi.get("MENU_Y")))

                if menu:
                    signal_coordinates[sig] = menu

                aspects = {}

                yellow = parse_point(cell(row, hi.get("YELLOW_X")), cell(row, hi.get("YELLOW_Y")))
                route_init = parse_point(cell(row, hi.get("ROUTEINIT_X")), cell(row, hi.get("ROUTEINIT_Y")))

                if yellow:
                    aspects["CALLING_ON"] = yellow
                if route_init:
                    aspects["INITIATION_INDICATION"] = route_init

                if aspects:
                    signal_aspects[sig] = aspects

                loaded_cal += 1

    # ---- SHUNT (suffix '9SH' -> prefix 'SH9') ----
    if "SHUNT" in wb_u.sheetnames:

        ws_u = wb_u["SHUNT"]
        hi = header_index_of(ws_u)
        sig_col = hi.get("SIGNAL")

        if sig_col is not None:

            for row in ws_u.iter_rows(min_row=2, values_only=True):

                if not row or cell(row, sig_col) is None:
                    continue

                raw_sig = str(cell(row, sig_col)).strip().upper()

                if raw_sig.endswith("SH"):
                    prefixed_sig = f"SH{raw_sig[:-2]}"
                elif raw_sig.startswith("SH"):
                    prefixed_sig = raw_sig
                else:
                    prefixed_sig = f"SH{raw_sig}"

                menu = parse_point(cell(row, hi.get("MENU_X")), cell(row, hi.get("MENU_Y")))

                if menu:
                    signal_coordinates[prefixed_sig] = menu

                indicator = parse_point(cell(row, hi.get("INDICATOR_X")), cell(row, hi.get("INDICATOR_Y")))
                route_init = parse_point(cell(row, hi.get("ROUTEINIT_X")), cell(row, hi.get("ROUTEINIT_Y")))
                snap_val = cell(row, hi.get("SNAPSHOT_PATH"))
                snapshot_path = str(snap_val).strip() if snap_val else None

                shunt_signals[prefixed_sig] = {}

                if indicator:

                    shunt_signals[prefixed_sig]["aspect_indicator"] = indicator

                    if snapshot_path:
                        shunt_signals[prefixed_sig]["aspect_snapshot"] = snapshot_path

                if route_init:

                    shunt_signals[prefixed_sig]["initiation_indicator"] = route_init
                    # No dedicated snapshot for Route Init in the
                    # universal file - shunt_signal_changed() checks
                    # this one by color (yellow = accepted) instead.

                loaded_shunt += 1

    log("")
    log("================================")
    log("UNIVERSAL YARD COORDINATES LOADED")
    log("================================")
    log(
        f"MAIN: {loaded_main}   "
        f"CALLING-ON: {loaded_cal}   "
        f"SHUNT: {loaded_shunt}"
    )
    if loaded_shunt:
        log(
            "NOTE: shunt Aspect Indicator uses ORB feature matching; "
            "Initiation Indication uses a color check (no dedicated "
            "snapshot for that point in the universal file)."
        )


def keyboard_pause_toggle(event=None):

    global paused

    paused = not paused

    if paused:

        status_label.config(
            text="PAUSED",
            fg="#f59e0b"
        )
        log("")
        log("================================")
        log("AUTOMATION PAUSED")
        log("================================")

    else:
        status_label.config(
            text="RUNNING",
            fg="#16a34a"
        )
        log("")
        log("================================")
        log("AUTOMATION RESUMED")
        log("================================")
def import_excel():
    upload_excel()
def start_run():
    start_automation()
def stop_run():
    stop_automation()
# =========================================================
# GUI
# =========================================================

root = tk.Tk()

root.title(
    "CROSS TABLE TESTING AUTOMATION"
)

root.geometry("1450x900")
root.configure(bg="#e9edf2")
root.state("zoomed")

# =========================================================
# STYLE
# =========================================================

style = ttk.Style()

style.theme_use("clam")

style.configure(
    "Treeview",
    background="white",
    foreground="black",
    rowheight=35,
    fieldbackground="white",
    font=("Segoe UI", 10)
)

style.configure(
    "Treeview.Heading",
    background="#1e293b",
    foreground="white",
    font=("Segoe UI", 11, "bold")
)

style.map(
    "Treeview",
    background=[("selected", "#2563eb")]
)

# =========================================================
# TOP HEADER
# =========================================================

header_frame = tk.Frame(
    root,
    bg="#0f172a",
    height=100
)

header_frame.pack(
    fill="x"
)

header_frame.pack_propagate(False)

# =========================================================
# LEFT LOGO
# =========================================================

left_logo_frame = tk.Frame(
    header_frame,
    bg="#0f172a"
)

left_logo_frame.pack(
    side="left",
    padx=20
)

try:

    railway_logo = tk.PhotoImage(
        file="indian_railways.png"
    )

    railway_logo = railway_logo.subsample(2, 2)

    tk.Label(
        left_logo_frame,
        image=railway_logo,
        bg="#0f172a"
    ).pack()

except:

    tk.Label(
        left_logo_frame,
        text="INDIAN RAILWAYS",
        font=("Segoe UI", 12, "bold"),
        bg="#0f172a",
        fg="white"
    ).pack()

# =========================================================
# TITLE
# =========================================================

title_frame = tk.Frame(
    header_frame,
    bg="#0f172a"
)

title_frame.pack(
    side="left",
    expand=True
)

tk.Label(

    title_frame,

    text="CROSS TABLE AUTOMATION SYSTEM",

    font=("Segoe UI", 24, "bold"),

    bg="#0f172a",

    fg="white"

).pack(
    pady=(15, 0)
)

tk.Label(

    title_frame,

    text="",

    font=("Segoe UI", 11),

    bg="#0f172a",

    fg="#cbd5e1"

).pack()

# =========================================================
# RIGHT LOGO
# =========================================================

right_logo_frame = tk.Frame(
    header_frame,
    bg="#0f172a"
)

right_logo_frame.pack(
    side="right",
    padx=20
)

try:

    company_logo = tk.PhotoImage(
        file="company_logo.png"
    )

    company_logo = company_logo.subsample(1, 1)

    tk.Label(
        right_logo_frame,
        image=company_logo,
        bg="#0f172a"
    ).pack()

except:

    tk.Label(
        right_logo_frame,
        text="COMPANY LOGO",
        font=("Segoe UI", 12, "bold"),
        bg="#0f172a",
        fg="white"
    ).pack()

# =========================================================
# MAIN AREA
# =========================================================

main_frame = tk.Frame(
    root,
    bg="#e9edf2"
)

main_frame.pack(
    fill="both",
    expand=True,
    padx=15,
    pady=15
)

# =========================================================
# LEFT PANEL
# =========================================================

left_panel = tk.Frame(
    main_frame,
    bg="white",
    bd=1,
    relief="solid"
)

left_panel.pack(
    side="left",
    fill="y",
    padx=(0, 10)
)

# =========================================================
# CONTROL TITLE
# =========================================================

tk.Label(

    left_panel,

    text="CONTROL PANEL",

    font=("Segoe UI", 15, "bold"),

    bg="white",

    fg="#0f172a"

).pack(
    pady=20
)

# =========================================================
# BUTTON STYLE
# =========================================================

def create_button(text, command, color):

    return tk.Button(

        left_panel,

        text=text,

        command=command,

        bg=color,

        fg="white",

        activebackground=color,

        activeforeground="white",

        relief="flat",

        cursor="hand2",

        font=("Segoe UI", 11, "bold"),

        width=26,

        height=2
    )

# =========================================================
# SIGNAL CONFIG SECTION
# =========================================================

signal_frame = tk.Frame(
    left_panel,
    bg="white"
)

signal_frame.pack(pady=10)

# =========================================================
# DUMMY MAIN BUTTON
# =========================================================

dummy_btn = tk.Button(

    signal_frame,

    text="CAPTURE SIGNALLING GEARS",

    bg="#2563eb",

    fg="white",

    activebackground="#2563eb",

    activeforeground="white",

    relief="flat",

    font=("Segoe UI", 11, "bold"),

    width=26,

    height=2
)

dummy_btn.pack(pady=(0, 8))

# =========================================================
# OTHER BUTTONS
# =========================================================

create_button(
    "CROSS TABLE",
    import_excel,
    "#2563eb"
).pack(pady=8)
excel_label = tk.Label(
    left_panel,
    text="No Excel Selected",
    bg="white",
    fg="blue",
    font=("Segoe UI", 9)
)

excel_label.pack(pady=(0,10))

create_button(
    "LOAD YARD COORDINATES",
    load_universal_coordinates,
    "#7c3aed"
).pack(pady=8)

create_button(
    "START TESTING",
    start_run,
    "#16a34a"
).pack(pady=20)

create_button(
    "STOP TESTING",
    stop_run,
    "#dc2626"
).pack(pady=8)

# =========================================================
# STATUS BOX
# =========================================================



# ADD STATUS FRAME HERE
status_frame = tk.Frame(
    left_panel,
    bg="#f8fafc",
    bd=1,
    relief="solid"
)

status_frame.pack(
    fill="x",
    padx=15,
    pady=25
)

tk.Label(
    status_frame,
    text="SYSTEM STATUS",
    font=("Segoe UI", 11, "bold"),
    bg="#f8fafc",
    fg="#0f172a"
).pack(pady=10)

status_label = tk.Label(
    status_frame,
    text="READY",
    font=("Segoe UI", 16, "bold"),
    bg="#f8fafc",
    fg="#16a34a"
)

status_label.pack(
    pady=(0, 15)
)
# =========================================================
# RIGHT PANEL
# =========================================================

right_panel = tk.Frame(
    main_frame,
    bg="#e9edf2"
)

right_panel.pack(
    side="left",
    fill="both",
    expand=True
)

# =========================================================
# TABLE TITLE
# =========================================================

table_title = tk.Label(

    right_panel,

    text="CROSS TABLE DETAILS",

    font=("Segoe UI", 16, "bold"),

    bg="#e9edf2",

    fg="#0f172a"

)

table_title.pack(
    anchor="w",
    pady=(0, 10)
)

# =========================================================
# TABLE FRAME
# =========================================================

table_frame = tk.Frame(
    right_panel,
    bg="white",
    bd=1,
    relief="solid"
)

table_frame.pack(
    fill="both",
    expand=True
)

# =========================================================
# SCROLLBAR
# =========================================================

tree_scroll = ttk.Scrollbar(
    table_frame
)

tree_scroll.pack(
    side="right",
    fill="y"
)

# =========================================================
# TREEVIEW
# =========================================================

tree = ttk.Treeview(
    table_frame,
    columns=(
        "NO",
        "SIGNAL",
        "ROUTES"
    ),
    show="headings",
    yscrollcommand=tree_scroll.set
)

tree.heading("NO", text="NO")
tree.heading("SIGNAL", text="SIGNAL")
tree.heading("ROUTES", text="TOTAL ROUTES")

tree.column("NO", width=80, anchor="center")
tree.column("SIGNAL", width=250, anchor="center")
tree.column("ROUTES", width=200, anchor="center")

tree.pack(fill="both", expand=True)

# =========================================================
# LOG TITLE
# =========================================================

log_title = tk.Label(

    right_panel,

    text="LIVE OPERATION LOG",

    font=("Segoe UI", 16, "bold"),

    bg="#e9edf2",

    fg="#0f172a"

)

log_title.pack(
    anchor="w",
    pady=(20, 10)
)

# =========================================================
# LOG FRAME
# =========================================================

log_frame = tk.Frame(
    right_panel,
    bg="white",
    bd=1,
    relief="solid"
)

log_frame.pack(
    fill="both",
    expand=True
)

# =========================================================
# LOG TEXT
# =========================================================

log_box = tk.Text(

    log_frame,

    bg="white",

    fg="black",

    font=("Consolas", 10),

    relief="flat"

)

log_box.pack(
    fill="both",
    expand=True,
    padx=10,
    pady=10
)


# =========================================================
# FOOTER
# =========================================================

footer = tk.Label(

    root,

    text="CROSS TABLE AUTOMATION SYSTEM",

    bg="#0f172a",

    fg="white",

    font=("Segoe UI", 10)

)

footer.pack(
    fill="x"
)

# =========================================================
# RUN
# =========================================================
keyboard.add_hotkey("p", keyboard_pause_toggle)
root.mainloop()
